import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os
import shutil

def build_viva_guide():
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    COLOR_TEAL = RGBColor(8, 125, 102)     # #087D66
    COLOR_DARK = RGBColor(30, 41, 59)     # Slate
    HEX_TEAL = "087D66"

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    r1 = p_title.add_run("MEDISAFE CHAIN")
    r1.font.name = "Arial"
    r1.font.size = Pt(24)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_TEAL

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(18)
    r2 = p_sub.add_run("Prescription Supply Chain Workflow Guide (Viva Demo Edition)")
    r2.font.name = "Arial"
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = COLOR_DARK

    def add_step_heading(title_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="14" w:space="4" w:color="{HEX_TEAL}"/></w:pBdr>')
        p._p.get_or_add_pPr().append(pBdr)
        
        r = p.add_run(title_text)
        r.font.name = "Arial"
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = COLOR_TEAL

    def add_item(num_text, detail_text):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Inches(0.1)
        r = p.add_run(f"{num_text}. {detail_text}")
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_DARK

    def add_field(label, value):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Inches(0.3)
        
        r_bullet = p.add_run("• ")
        r_bullet.font.name = "Arial"
        r_bullet.font.size = Pt(10)
        r_bullet.font.color.rgb = COLOR_TEAL

        r_lbl = p.add_run(f"{label}: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_DARK

        r_val = p.add_run(value)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10)
        r_val.font.color.rgb = COLOR_DARK

    def add_result(text, is_final=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.left_indent = Inches(0.1)

        r_lbl = p.add_run("Result: ")
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_TEAL if is_final else COLOR_DARK

        r_val = p.add_run(text)
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10)
        r_val.font.italic = not is_final
        r_val.font.bold = is_final
        r_val.font.color.rgb = COLOR_TEAL if is_final else COLOR_DARK

    # STEP 1
    add_step_heading("STEP 1: Patient Prepares Credential")
    add_item("1", 'Open http://localhost:5173 (or http://172.28.0.60:5173) and click "Prepare a patient credential".')
    add_item("2", 'Enter these exact Viva demo values:')
    add_field("National ID / Passport", "199854120330")
    add_field("Full legal name", "Kamal Perera")
    add_field("Date of birth", "05/12/1998")
    add_field("Credential passphrase", "VivaDemoPass2026!")
    add_item("3", 'Click "Create & download credential".')
    add_result("A file named medisafe-patient-preparation.json is saved in Downloads. Copy the public Patient Hash (0x...) shown on screen.")

    # STEP 2
    add_step_heading("STEP 2: Doctor Issues Prescription on Blockchain")
    add_item("1", "In MetaMask, select the Doctor account (0x2a13ee1f4C69386864A271Bceb401fcd3d91c942).")
    add_item("2", "Sign in as Doctor on http://localhost:5173.")
    add_item("3", 'Under "Issue a prescription", enter:')
    add_field("Patient privacy hash", "[Paste the Patient Hash copied in Step 1]")
    add_field("Medicine ID", "MED-AUGMENTIN-625")
    add_field("Batch ID", "BATCH-VIVA-001  (or BATCH-SAFE-001)")
    add_field("Valid for", "30 days")
    add_item("4", 'Click "Issue prescription" and click Confirm in MetaMask.')
    add_result("Copy the Prescription ID (0x...) displayed in the green success notification.")

    # STEP 3
    add_step_heading("STEP 3: Patient Verifies & Prepares Pharmacy File")
    add_item("1", "In MetaMask, switch to the Patient account (0x2376C6970c55AB66c4FFcad53A49fb21f862B613).")
    add_item("2", "Sign in as Patient using the Prescription ID copied from Step 2.")
    add_item("3", 'On the "Verify your identity" screen, enter:')
    add_field("Encrypted preparation credential", "Upload the latest medisafe-patient-preparation.json from Downloads")
    add_field("Credential passphrase", "VivaDemoPass2026!")
    add_field("New pharmacy presentation passphrase", "PharmacyPass2026!")
    add_field("Intended pharmacy wallet", "0x7465C8b4C7a9068c440044222539f993bf4537A7")
    add_item("4", 'Click "Verify and prepare pharmacy handoff".')
    add_result("A new file named medisafe-pharmacy-presentation.json is downloaded automatically.")

    # STEP 4
    add_step_heading("STEP 4: Pharmacy Dispenses Medication")
    add_item("1", "In MetaMask, switch to the Pharmacy account (0x7465C8b4C7a9068c440044222539f993bf4537A7).")
    add_item("2", "Sign in as Pharmacy on http://localhost:5173.")
    add_item("3", 'On the Pharmacy desk, enter:')
    add_field("Encrypted patient presentation file", "Upload medisafe-pharmacy-presentation.json from Downloads")
    add_field("Credential passphrase", "PharmacyPass2026!")
    add_item("4", 'Click "Verify & dispense" and click Confirm in MetaMask.')
    add_result("The prescription is dispensed and officially marked as USED on the blockchain!", is_final=True)

    # Save
    local_name = "Medisafe_Viva_Workflow_Guide.docx"
    doc.save(local_name)

    downloads_dir = os.path.expanduser("~/Downloads")
    dest_name = os.path.join(downloads_dir, "Medisafe_Viva_Workflow_Guide.docx")
    shutil.copy2(local_name, dest_name)

    try:
        shutil.copy2(local_name, os.path.join(downloads_dir, "Medisafe_Workflow_Guide_Simple.docx"))
        shutil.copy2(local_name, os.path.join(downloads_dir, "Medisafe_Workflow_Guide.docx"))
    except Exception:
        pass

    print("SUCCESS: Viva guide generated and copied to Downloads!")

if __name__ == "__main__":
    build_viva_guide()
