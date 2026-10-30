import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import shutil
import os

def create_presentation_guide():
    doc = docx.Document()

    # Set Margins (0.8 inch for clean layout)
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette - Professional Medical Teal & Navy
    C_PRIMARY = RGBColor(8, 125, 102)     # #087D66 Deep Teal
    C_DARK = RGBColor(15, 23, 42)         # #0F172A Navy Slate
    C_MUTED = RGBColor(71, 85, 105)       # #475569 Slate Muted
    C_ACCENT = RGBColor(14, 116, 144)     # #0E7490 Cyan Accent
    C_ORANGE = RGBColor(194, 65, 12)      # Warm callout

    HEX_TEAL = "087D66"
    HEX_BG_LIGHT = "F0FDF4"               # Soft Mint
    HEX_BG_GRAY = "F8FAFC"                # Slate 50
    HEX_BORDER = "CBD5E1"                 # Slate 300

    def set_cell_background(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
            node = OxmlElement(m)
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # 1. HEADER BANNER
    header_tbl = doc.add_table(rows=1, cols=1)
    header_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = header_tbl.rows[0].cells[0]
    c.width = Inches(6.9)
    set_cell_background(c, "087D66")
    set_cell_margins(c, 160, 160, 200, 200)

    p0 = c.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r0 = p0.add_run("MEDISAFE CHAIN — COMPONENT 4\n")
    r0.font.name = "Arial"
    r0.font.size = Pt(11)
    r0.font.bold = True
    r0.font.color.rgb = RGBColor(204, 251, 241) # Light Mint

    r1 = p0.add_run("Live Presentation & Demo Cheat-Sheet\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(20)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(255, 255, 255)

    r2 = p0.add_run("Prescription Management: Privacy • Blockchain • Multi-Layer Validation")
    r2.font.name = "Arial"
    r2.font.size = Pt(10)
    r2.font.color.rgb = RGBColor(240, 253, 250)

    doc.add_paragraph() # Spacing

    # 2. PRE-DEMO CHECKLIST
    p_chk = doc.add_paragraph()
    r = p_chk.add_run("⚡ Pre-Demo Checklist (Have these open before starting)")
    r.font.name = "Arial"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = C_PRIMARY

    chk_tbl = doc.add_table(rows=3, cols=2)
    chk_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    chk_items = [
        ("Terminal 1 (Blockchain)", "Running: npx hardhat node (Port 8545)"),
        ("Terminal 2 (API Backend)", "Running: node server.js (Port 5004)"),
        ("Browser Window", "Open: http://localhost:5173 (MetaMask on Hardhat 31337)")
    ]
    for idx, (k, v) in enumerate(chk_items):
        row = chk_tbl.rows[idx]
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.7)
        set_cell_background(row.cells[0], "F1F5F9")
        set_cell_background(row.cells[1], "F8FAFC")
        set_cell_margins(row.cells[0], 60, 60, 100, 100)
        set_cell_margins(row.cells[1], 60, 60, 100, 100)
        
        p = row.cells[0].paragraphs[0]
        run = p.add_run(k)
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = C_DARK

        p2 = row.cells[1].paragraphs[0]
        run2 = p2.add_run(v)
        run2.font.name = "Arial"
        run2.font.size = Pt(9.5)
        run2.font.color.rgb = C_DARK

    doc.add_paragraph()

    # HELPER: STEP BOX
    def add_presentation_step(step_num, step_title, actor_badge, actions_list, inputs_dict, talking_points):
        # Step Header Box
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.9)
        set_cell_background(cell, "087D66")
        set_cell_margins(cell, 90, 90, 140, 140)

        p = cell.paragraphs[0]
        r_step = p.add_run(f"STEP {step_num}: {step_title.upper()}   ")
        r_step.font.name = "Arial"
        r_step.font.size = Pt(11)
        r_step.font.bold = True
        r_step.font.color.rgb = RGBColor(255, 255, 255)

        r_badge = p.add_run(f" [{actor_badge}] ")
        r_badge.font.name = "Arial"
        r_badge.font.size = Pt(9.5)
        r_badge.font.bold = True
        r_badge.font.color.rgb = RGBColor(204, 251, 241)

        # Actions
        doc.add_paragraph()
        p_act = doc.add_paragraph()
        r_act_title = p_act.add_run("👉 What to Click & Do:")
        r_act_title.font.name = "Arial"
        r_act_title.font.size = Pt(10)
        r_act_title.font.bold = True
        r_act_title.font.color.rgb = C_PRIMARY

        for act in actions_list:
            p_a = doc.add_paragraph()
            p_a.paragraph_format.left_indent = Inches(0.2)
            p_a.paragraph_format.space_after = Pt(2)
            ra = p_a.add_run(f"• {act}")
            ra.font.name = "Arial"
            ra.font.size = Pt(9.5)
            ra.font.color.rgb = C_DARK

        # Data Input Table
        if inputs_dict:
            doc.add_paragraph()
            in_tbl = doc.add_table(rows=len(inputs_dict) + 1, cols=2)
            in_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            in_tbl.rows[0].cells[0].width = Inches(2.2)
            in_tbl.rows[0].cells[1].width = Inches(4.7)
            set_cell_background(in_tbl.rows[0].cells[0], "E2E8F0")
            set_cell_background(in_tbl.rows[0].cells[1], "E2E8F0")
            set_cell_margins(in_tbl.rows[0].cells[0], 50, 50, 80, 80)
            set_cell_margins(in_tbl.rows[0].cells[1], 50, 50, 80, 80)
            
            p_h1 = in_tbl.rows[0].cells[0].paragraphs[0]
            rh1 = p_h1.add_run("Field Name")
            rh1.font.bold = True
            rh1.font.size = Pt(9)
            
            p_h2 = in_tbl.rows[0].cells[1].paragraphs[0]
            rh2 = p_h2.add_run("Exact Demo Value to Enter / Click")
            rh2.font.bold = True
            rh2.font.size = Pt(9)

            for i, (field, val) in enumerate(inputs_dict.items()):
                row = in_tbl.rows[i + 1]
                row.cells[0].width = Inches(2.2)
                row.cells[1].width = Inches(4.7)
                set_cell_background(row.cells[0], "F8FAFC")
                set_cell_background(row.cells[1], "FFFFFF")
                set_cell_margins(row.cells[0], 50, 50, 80, 80)
                set_cell_margins(row.cells[1], 50, 50, 80, 80)

                pf = row.cells[0].paragraphs[0]
                rf = pf.add_run(field)
                rf.font.bold = True
                rf.font.size = Pt(8.5)
                rf.font.color.rgb = C_DARK

                pv = row.cells[1].paragraphs[0]
                rv = pv.add_run(val)
                rv.font.size = Pt(8.5)
                rv.font.color.rgb = C_DARK

        # Talking Points Box
        doc.add_paragraph()
        speech_tbl = doc.add_table(rows=1, cols=1)
        speech_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        scell = speech_tbl.rows[0].cells[0]
        scell.width = Inches(6.9)
        set_cell_background(scell, "F0FDF4") # soft green
        set_cell_margins(scell, 80, 80, 120, 120)

        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:color="087D66"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        scell._tc.get_or_add_tcPr().append(borders)

        sp = scell.paragraphs[0]
        rsp = sp.add_run("🎤 WHAT TO SAY TO EVALUATORS:\n")
        rsp.font.name = "Arial"
        rsp.font.size = Pt(9.5)
        rsp.font.bold = True
        rsp.font.color.rgb = C_PRIMARY

        rsp_text = sp.add_run(f'"{talking_points}"')
        rsp_text.font.name = "Arial"
        rsp_text.font.size = Pt(9.5)
        rsp_text.font.italic = True
        rsp_text.font.color.rgb = C_DARK

        doc.add_paragraph() # Spacing between steps

    # --- STEP 1 ---
    add_presentation_step(
        1,
        "Patient Prepares Privacy Credential",
        "Actor: Patient (Self-Sovereign)",
        [
            "Open http://localhost:5173 on Chrome.",
            "Click 'Prepare a patient credential' on the bottom of the login screen.",
            "Click the 'Auto-fill test data' button (or type the values below).",
            "Click 'Create & download credential'.",
            "COPY the displayed green Patient Hash (0x...) to your clipboard!"
        ],
        {
            "National ID / Passport": "200112345678",
            "Full Legal Name": "John Doe",
            "Date of Birth": "07/21/2001",
            "Passphrase": "MedisafeSecret2026!",
            "Generated File": "medisafe-patient-preparation.json (saved in Downloads)"
        },
        "Notice that no patient PII—no name, date of birth, or National ID—ever leaves the patient's device. "
        "The browser derives a one-way Keccak-256 hash using a 32-byte salt and encrypts the file locally with AES-256-GCM. "
        "The patient only gives this public cryptographic hash to their doctor."
    )

    # --- STEP 2 ---
    add_presentation_step(
        2,
        "Doctor Issues Prescription on Blockchain",
        "Actor: Doctor (MetaMask 0x2a13e...)",
        [
            "In MetaMask, select Doctor Account (0x2a13ee1f4C69386864A271Bceb401fcd3d91c942).",
            "Select 'Doctor' role and click 'Connect MetaMask & Sign In'.",
            "Sign the EIP-191 authentication message in MetaMask.",
            "In 'Issue a prescription', paste the Patient Hash from Step 1.",
            "Fill Medicine ID and Batch ID as shown below.",
            "Click 'Issue prescription' and click Confirm in MetaMask.",
            "COPY the Prescription ID (0x...) shown in the green success notification!"
        ],
        {
            "Patient Privacy Hash": "[Paste the 0x... hash copied from Step 1]",
            "Medicine ID": "MED-AMOX-500",
            "Batch ID": "BATCH-SAFE-001  (or BATCH-2026-01)",
            "Valid For": "30 days",
            "On-Chain Output": "PrescriptionRegistry.issuePrescription() -> Status: ACTIVE"
        },
        "The doctor has now issued an authentic prescription directly onto the Ethereum blockchain. "
        "The contract verifies the doctor's verified DOCTOR_ROLE, binds the prescription to the patient hash and medicine batch, "
        "and emits an immutable PrescriptionIssued event without exposing who the patient is."
    )

    # --- STEP 3 ---
    add_presentation_step(
        3,
        "Patient Creates Pharmacy-Scoped Presentation",
        "Actor: Patient (MetaMask 0x2376C...)",
        [
            "In MetaMask, switch to Patient Account (0x2376C6970c55AB66c4FFcad53A49fb21f862B613).",
            "Select 'Patient' role and paste the Prescription ID from Step 2.",
            "Click 'Connect MetaMask & Sign In' and sign the challenge.",
            "On the 'Verify your identity' screen, upload the downloaded medisafe-patient-preparation file.",
            "Enter passphrases and target pharmacy wallet as shown below.",
            "Click 'Verify and prepare pharmacy handoff'.",
            "Result: A new file 'medisafe-pharmacy-presentation.json' is downloaded automatically."
        ],
        {
            "Encrypted Credential": "Upload medisafe-patient-preparation.json from Downloads",
            "Credential Passphrase": "MedisafeSecret2026!",
            "New Presentation Passphrase": "PharmacyPass123!",
            "Intended Pharmacy Wallet": "0x7465C8b4C7a9068c440044222539f993bf4537A7",
            "Generated File": "medisafe-pharmacy-presentation.json"
        },
        "This step prevents replay attacks and medical identity theft. "
        "The patient generates a scoped presentation that is cryptographically tied ONLY to this specific pharmacy's wallet address and this prescription. "
        "Even if an attacker steals this file, no other pharmacy can dispense it."
    )

    # --- STEP 4 ---
    add_presentation_step(
        4,
        "Pharmacy Verifies & Dispenses Medication",
        "Actor: Pharmacy (MetaMask 0x7465C...)",
        [
            "In MetaMask, switch to Pharmacy Account (0x7465C8b4C7a9068c440044222539f993bf4537A7).",
            "Select 'Pharmacy' role and sign in with MetaMask.",
            "On the Pharmacy Desk, upload the medisafe-pharmacy-presentation.json file.",
            "Enter Credential Passphrase: PharmacyPass123!",
            "Click 'Verify & dispense medication'.",
            "Click Confirm in MetaMask.",
            "Result: SUCCESS! The prescription is officially dispensed and marked USED on blockchain."
        ],
        {
            "Presentation File": "Upload medisafe-pharmacy-presentation.json",
            "Credential Passphrase": "PharmacyPass123!",
            "On-Chain Execution": "verifyAndDispenseWithPatientHash(rxId, presentedPatientHash)",
            "Final State": "PrescriptionStatus.USED (Permanently locked against reuse)"
        },
        "This is the core strength of our architecture. In a single atomic smart contract transaction: "
        "First, it verifies the patient's hash matches. "
        "Second, it queries Component 2 AI model to confirm the medicine batch is safe (risk score <= 69). "
        "Third, it queries Component 3 to confirm this pharmacy is trusted (score >= 70). "
        "Finally, it locks the prescription to USED so it can NEVER be reused or double-dispensed!"
    )

    # SECTION 4: CHEATSHEET TABLE
    p_sum = doc.add_paragraph()
    r_s = p_sum.add_run("📋 Quick-Copy Cheatsheet for Presentation")
    r_s.font.name = "Arial"
    r_s.font.size = Pt(12)
    r_s.font.bold = True
    r_s.font.color.rgb = C_PRIMARY

    sum_tbl = doc.add_table(rows=6, cols=3)
    sum_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_headers = ["Item / Step", "Value / Address", "Role / Notes"]
    for idx, text in enumerate(s_headers):
        c = sum_tbl.rows[0].cells[idx]
        set_cell_background(c, "087D66")
        set_cell_margins(c, 70, 70, 90, 90)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.bold = True
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    sum_rows = [
        ("Doctor Account", "0x2a13ee1f4C69386864A271Bceb401fcd3d91c942", "DOCTOR_ROLE (MetaMask Account 1)"),
        ("Patient Account", "0x2376C6970c55AB66c4FFcad53A49fb21f862B613", "Patient Identity (MetaMask Account 2)"),
        ("Pharmacy Account", "0x7465C8b4C7a9068c440044222539f993bf4537A7", "PHARMACY_ROLE (MetaMask Account 3)"),
        ("Patient Passphrase", "MedisafeSecret2026!", "Used to encrypt patient credential file"),
        ("Pharmacy Passphrase", "PharmacyPass123!", "Used to encrypt pharmacy presentation file")
    ]
    for row_idx, data in enumerate(sum_rows):
        row = sum_tbl.rows[row_idx + 1]
        for col_idx, text in enumerate(data):
            c = row.cells[col_idx]
            set_cell_background(c, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(c, 50, 50, 70, 70)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            r.font.color.rgb = C_DARK

    doc.add_paragraph()

    # SECTION 5: TOP 3 EVALUATOR QUESTIONS
    p_q = doc.add_paragraph()
    rq = p_q.add_run("💡 Answers to Top Evaluator Questions")
    rq.font.name = "Arial"
    rq.font.size = Pt(12)
    rq.font.bold = True
    rq.font.color.rgb = C_PRIMARY

    qa_list = [
        ("Q1: How do you protect patient privacy if the blockchain is public?",
         "We never store patient names, DOBs, or NICs on-chain. Only a salt-randomized Keccak-256 hash is anchored. Identity credentials are encrypted client-side using AES-256-GCM and remain in the patient's local custody."),
        ("Q2: What prevents double-dispensing or counterfeit reuse?",
         "The smart contract enforces a strict state machine. The moment verifyAndDispenseWithPatientHash() is called, status transitions to USED in the same atomic block. Any subsequent attempt immediately reverts with PrescriptionNotActive."),
        ("Q3: How does Component 4 connect with the other components?",
         "Component 4 interfaces with Component 2 (AI Risk Scoring) to guarantee the medicine batch risk score is safe (<= 69), and Component 3 (Pharmacy Trust) to guarantee the dispensing pharmacy holds a verified trust rating (>= 70).")
    ]
    for q, a in qa_list:
        p_item = doc.add_paragraph()
        p_item.paragraph_format.space_after = Pt(4)
        rq_title = p_item.add_run(f"• {q}\n")
        rq_title.font.name = "Arial"
        rq_title.font.bold = True
        rq_title.font.size = Pt(9.5)
        rq_title.font.color.rgb = C_PRIMARY

        ra_text = p_item.add_run(f"  {a}")
        ra_text.font.name = "Arial"
        ra_text.font.size = Pt(9)
        ra_text.font.color.rgb = C_DARK

    filename = "Component_4_Presentation_Guide.docx"
    doc.save(filename)
    print(f"SUCCESS: {filename} generated!")

    # Copy to Downloads
    user_downloads = os.path.expanduser("~/Downloads")
    dest_path = os.path.join(user_downloads, filename)
    shutil.copy2(filename, dest_path)
    print(f"SUCCESS: Copied to {dest_path}")

if __name__ == "__main__":
    create_presentation_guide()
