from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs/rp-diary/Component_2_RP_Diary_October_2026.docx")

NAVY = "173A4A"
TEAL = "2A8C7B"
PALE = "EEF6F4"
LIGHT = "F5F7F8"
BORDER = "D9D9D9"
TEXT = RGBColor(31, 46, 54)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, name="Aptos", size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def add_field_line(table, row_idx, label, value, col_idx):
    cell = table.cell(row_idx, col_idx)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(f"{label}  ")
    set_font(r, size=9.5, bold=True, color=RGBColor(76, 92, 99))
    r = p.add_run(value)
    set_font(r, size=10.5, bold=False)


def add_log_page(doc, period, entries, period_note=None):
    section = doc.sections[-1]
    section.top_margin = Inches(0.52)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.62)
    section.right_margin = Inches(0.62)

    title = doc.add_paragraph(style="Title")
    remove_paragraph_borders(title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    run = title.add_run("Research Project Periodic Log Entry Form")
    set_font(run, size=16, bold=True, color=RGBColor(0, 0, 0))

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(9)
    run = subtitle.add_run(period)
    set_font(run, size=11.5, bold=True, color=RGBColor(42, 140, 123))
    if period_note:
        run = subtitle.add_run(f"  {period_note}")
        set_font(run, size=9.5, color=RGBColor(96, 111, 118))

    info = doc.add_table(rows=2, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.autofit = False
    info.columns[0].width = Inches(3.55)
    info.columns[1].width = Inches(3.55)
    set_table_borders(info, color="E5E9EB", size="5")
    for row in info.rows:
        for cell in row.cells:
            set_cell_margins(cell, top=85, bottom=85)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    add_field_line(info, 0, "Group ID", "J26-IT-317", 0)
    add_field_line(info, 0, "Student ID", "____________________", 1)
    add_field_line(info, 1, "Name", "________________________________", 0)
    add_field_line(info, 1, "Individual Component", "Component 2  Medicine Risk Scoring", 1)

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.space_before = Pt(9)
    heading.paragraph_format.space_after = Pt(6)
    run = heading.add_run("Summary of RP work carried out during the reporting period")
    set_font(run, size=11.5, bold=True, color=RGBColor(23, 58, 74))

    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.45)
    table.columns[1].width = Inches(5.65)
    set_table_borders(table)
    hdr = table.rows[0].cells
    for cell, text_value in zip(hdr, ("Date", "RP Work")):
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=110, bottom=110)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text_value)
        set_font(r, size=10.5, bold=True, color=RGBColor(255, 255, 255))
    set_repeat_header(table.rows[0])

    for idx, (date, work) in enumerate(entries):
        cells = table.add_row().cells
        if idx % 2:
            for cell in cells:
                set_cell_shading(cell, LIGHT)
        for cell in cells:
            set_cell_margins(cell, top=95, bottom=95)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cells[0].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(date)
        set_font(r, size=9.6, bold=True, color=RGBColor(42, 97, 91))
        p = cells[1].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(work)
        set_font(r, size=9.6)

    sup = doc.add_paragraph()
    sup.paragraph_format.space_before = Pt(8)
    sup.paragraph_format.space_after = Pt(3)
    r = sup.add_run("Supervisor Comment")
    set_font(r, size=10.5, bold=True, color=RGBColor(23, 58, 74))

    for _ in range(2):
        p = doc.add_paragraph("________________________________________________________________________________")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 0.85
        for r in p.runs:
            set_font(r, size=8.5, color=RGBColor(156, 165, 169))

    progress = doc.add_paragraph()
    progress.paragraph_format.space_before = Pt(5)
    progress.paragraph_format.space_after = Pt(6)
    r = progress.add_run("Student progressing well in research project       Yes  /  No")
    set_font(r, size=10, bold=True)

    sign = doc.add_table(rows=1, cols=2)
    sign.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign.autofit = False
    sign.columns[0].width = Inches(4.6)
    sign.columns[1].width = Inches(2.5)
    for idx, label in enumerate(("Supervisor Signature", "Date")):
        cell = sign.cell(0, idx)
        set_cell_margins(cell, top=55, bottom=0)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("____________________________\n" + label)
        set_font(r, size=9.5, color=RGBColor(76, 92, 99))
    # Hide the signature layout table borders.
    set_table_borders(sign, color="FFFFFF", size="0")


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    title_p_pr = styles["Title"]._element.get_or_add_pPr()
    title_p_bdr = title_p_pr.find(qn("w:pBdr"))
    if title_p_bdr is not None:
        title_p_pr.remove(title_p_bdr)

    first_period = [
        ("1 Oct 2026", "Reviewed the Component 2 scope and identified the medicine batch data required for risk scoring."),
        ("4 Oct 2026", "Collected medicine data and checked the dataset for missing, duplicated and inconsistent values."),
        ("7 Oct 2026", "Created the initial project structure and finalized the workflow for training, inference and API access."),
        ("9 Oct 2026", "Cleaned the dataset and prepared model features, including medicine category, dosage form, strength and classification."),
        ("12 Oct 2026", "Developed Random Forest and XGBoost risk models, added prediction logic and generated SHAP explanations."),
        ("13 Oct 2026", "Compared model results, selected XGBoost as the preferred model and documented the evaluation findings."),
        ("14 Oct 2026", "Tested the risk scoring API and implemented blockchain transaction support for the MedicineRiskRegistry contract."),
    ]
    add_log_page(doc, "1 October to 14 October 2026", first_period)

    doc.add_section(WD_SECTION.NEW_PAGE)
    second_period = [
        ("15 Oct 2026", "Added the local risk-registry deployment workflow, centralized blockchain settings and connected registered traceability batches to risk assessment."),
        ("16 Oct 2026", "Checked the Component 2 and Component 4 integration rule so unassessed medicine batches are rejected before prescription processing."),
        ("18 Oct 2026", "Coordinated local Hardhat contract addresses and verified the expected risk-registry connection for the demonstration environment."),
        ("19 Oct 2026", "Added risk-record publishing and verification, recorded the local registry deployment and documented the Component 2 to Component 4 handoff."),
        ("20 Oct 2026", "Reviewed the progress of the other components and checked the end-to-end dependencies for traceability, trust scoring and prescription management."),
        ("21 Oct 2026", "Planned the shared Supabase PostgreSQL database structure and reviewed the Component 2 frontend for a more professional demonstration design."),
    ]
    add_log_page(doc, "15 October to 28 October 2026", second_period, "Current period through 21 October")

    props = doc.core_properties
    props.title = "Component 2 RP Diary October 2026"
    props.subject = "Medicine Risk Scoring periodic research log"
    props.author = "MEDISAFE CHAIN Component 2 Student"
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
