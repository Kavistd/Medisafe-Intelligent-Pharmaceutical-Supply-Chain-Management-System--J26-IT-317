from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path("outputs/rp-diary/Component_2_AI_Use_Record_2026.docx")
NAVY = "173A4A"
TEAL = RGBColor(42, 140, 123)
TEXT = RGBColor(31, 46, 54)
GRAY = RGBColor(86, 101, 108)


def font(run, size=10, bold=False, color=TEXT):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.find(qn("w:shd")) or OxmlElement("w:shd")
    if node.getparent() is None:
        pr.append(node)
    node.set(qn("w:fill"), fill)


def margins(cell, value=100):
    pr = cell._tc.get_or_add_tcPr()
    tc_mar = pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        pr.append(tc_mar)
    for name in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def borders(table):
    pr = table._tbl.tblPr
    block = pr.first_child_found_in("w:tblBorders")
    if block is None:
        block = OxmlElement("w:tblBorders")
        pr.append(block)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "D9D9D9")
        block.append(node)


def remove_title_border(style, paragraph):
    for ppr in (style._element.get_or_add_pPr(), paragraph._p.get_or_add_pPr()):
        node = ppr.find(qn("w:pBdr"))
        if node is not None:
            ppr.remove(node)


def add_labeled_paragraph(doc, label, value):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(label + "  ")
    font(r, bold=True, color=GRAY)
    r = p.add_run(value)
    font(r)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10)

    title = doc.add_paragraph(style="Title")
    remove_title_border(doc.styles["Title"], title)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("Artificial Intelligence Use Record")
    font(r, size=17, bold=True, color=RGBColor(0, 0, 0))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run("MEDISAFE CHAIN  Component 2 Medicine Risk Scoring")
    font(r, size=11.5, bold=True, color=TEAL)

    add_labeled_paragraph(doc, "Group ID", "J26-IT-317")
    add_labeled_paragraph(doc, "Student ID", "____________________")
    add_labeled_paragraph(doc, "Student Name", "________________________________________")
    add_labeled_paragraph(doc, "AI tools", "ChatGPT and Codex")
    add_labeled_paragraph(doc, "Overall use", "Supporting assistance under student review and control")

    statement = doc.add_paragraph()
    statement.paragraph_format.space_before = Pt(6)
    statement.paragraph_format.space_after = Pt(9)
    statement.paragraph_format.line_spacing = 1.08
    r = statement.add_run(
        "AI was used to support selected planning, explanation, troubleshooting, code review and writing tasks. "
        "I reviewed the suggestions, adapted them to the project, ran the implementation, checked the results and "
        "remained responsible for the final technical and academic work. AI output was not accepted without verification."
    )
    font(r, size=10.2)

    entries = [
        ("17-24 Sep 2026", "Project planning", "Helped organize Component 2 requirements and possible technologies.", "Selected the scope after team and supervisor discussions."),
        ("1-9 Oct 2026", "Data preparation", "Suggested checks for missing values, duplicates and feature preparation.", "Inspected the dataset and applied the relevant cleaning steps."),
        ("10-13 Oct 2026", "Model development", "Provided examples and explanations for Random Forest, XGBoost and SHAP.", "Trained, compared and evaluated the models using project data."),
        ("12-14 Oct 2026", "Risk API and tests", "Assisted with API structure, validation ideas and test-case review.", "Reviewed the code, ran the service and checked endpoint results."),
        ("14-19 Oct 2026", "Blockchain integration", "Helped explain deployment, configuration and transaction verification steps.", "Ran the local Hardhat commands and verified addresses and responses."),
        ("15-20 Oct 2026", "Component integration", "Supported analysis of Component 1 and Component 4 data flows and error messages.", "Confirmed the interfaces with team members and tested available flows."),
        ("19-21 Oct 2026", "Frontend review", "Suggested interface wording, layout improvements and troubleshooting steps.", "Reviewed the interface and tested it in the local environment."),
        ("21 Oct 2026", "Database planning", "Compared database options and suggested a shared Supabase PostgreSQL structure.", "Chose the database approach with the team and retained responsibility for configuration."),
        ("Throughout", "Documentation", "Helped improve grammar, structure and clarity in technical notes and PR descriptions.", "Checked every statement against the implemented project before use."),
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [1.05, 1.25, 2.45, 2.55]
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    borders(table)
    headers = ("Period", "Project activity", "AI assistance", "Student work and verification")
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, NAVY)
        margins(cell, 95)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        font(r, size=9, bold=True, color=RGBColor(255, 255, 255))

    for row_index, row in enumerate(entries):
        cells = table.add_row().cells
        if row_index % 2:
            for cell in cells:
                shade(cell, "F3F6F7")
        for idx, (cell, value) in enumerate(zip(cells, row)):
            margins(cell, 80)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            if idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            font(r, size=8.5, bold=(idx == 0), color=TEAL if idx == 0 else TEXT)

    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(9)
    h.paragraph_format.space_after = Pt(3)
    r = h.add_run("Academic integrity declaration")
    font(r, size=11, bold=True, color=RGBColor(23, 58, 74))

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.05
    r = p.add_run(
        "I confirm that this record accurately describes my use of AI assistance. I did not use AI to fabricate research "
        "data, test results, references or supervisor feedback. I will follow the institution's disclosure and citation "
        "requirements and provide prompts or outputs if requested."
    )
    font(r, size=9.5)

    sig = doc.add_table(rows=1, cols=2)
    sig.autofit = False
    sig.columns[0].width = Inches(4.6)
    sig.columns[1].width = Inches(2.6)
    for cell, label in zip(sig.rows[0].cells, ("Student Signature", "Date")):
        p = cell.paragraphs[0]
        r = p.add_run("____________________________\n" + label)
        font(r, size=9.3, color=GRAY)

    props = doc.core_properties
    props.title = "Component 2 Artificial Intelligence Use Record"
    props.subject = "Disclosure of AI assistance in the MEDISAFE CHAIN research project"
    props.author = "MEDISAFE CHAIN Component 2 Student"
    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    build()
